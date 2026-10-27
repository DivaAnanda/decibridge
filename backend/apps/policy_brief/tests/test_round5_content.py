"""Round 5 item 2: the issued brief carries the PSA, the decision question, the
target population and the KFT's final decision.

Pak Anom: "Brief: tambahkan hasil PSA + pertanyaan keputusan + populasi target ke
PDF/DOCX." The decision of item 3 goes in as well, since the brief is the record a
reader takes away from the committee.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from apps.econ.models import EconomicModel
from apps.econ.tests.factories import make_psa_result
from apps.policy_brief import service
from apps.policy_brief.models import PolicyBriefDocument
from apps.recommendation.models import FinalDecisionType
from apps.recommendation.tests.factories import make_final_decision

pytestmark = pytest.mark.django_db


def _text(brief: PolicyBriefDocument) -> str:
    doc = Document(io.BytesIO(service.absolute_path(brief.docx_path).read_bytes()))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


@pytest.fixture
def full_case(approved_case_with_rec, hta_user, kft_member_user):
    case, rec = approved_case_with_rec
    make_psa_result(case, hta_user)
    make_final_decision(
        case,
        kft_member_user,
        decision=FinalDecisionType.REASSESS_AFTER_NEGOTIATION,
        recommendation=rec,
    )
    return case


def test_generated_brief_carries_the_round_5_content(hta_client, full_case):
    response = hta_client.post(f"/api/v1/cases/{full_case.case_id}/policy-briefs/")
    assert response.status_code == 201, response.data

    text = _text(PolicyBriefDocument.objects.get(pk=response.data["id"]))

    assert "Apakah ARNI perlu masuk formularium?" in text
    assert "Pasien HFrEF dewasa" in text
    assert full_case.population in text
    assert "Analisis Sensitivitas Probabilistik (PSA)" in text
    assert "56.0%" in text
    assert "Kaji ulang setelah negosiasi harga (Reassess after price negotiation)" in text
    assert "Alasan pengujian." in text


class TestContextBuilders:
    def test_psa_on_the_current_inputs_is_current(self, full_case):
        assert service.build_context(full_case, version=1).psa.is_current is True

    def test_psa_with_changed_inputs_is_not(self, full_case):
        model = EconomicModel.objects.get(case=full_case)
        model.wtp_threshold = 1
        model.save()

        assert service.build_context(full_case, version=1).psa.is_current is False

    def test_target_population_falls_back_to_the_pico(self, full_case):
        type(full_case).objects.filter(pk=full_case.pk).update(population="")
        full_case.refresh_from_db()

        ctx = service.build_context(full_case, version=1)

        assert ctx.case.population == "Pasien HFrEF dewasa"

    def test_no_final_decision_leaves_it_empty(self, approved_case_with_rec):
        case, _ = approved_case_with_rec

        assert service.build_context(case, version=1).final_decision is None
