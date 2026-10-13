"""Builds the manual book (DOCX and PDF) from one captured walkthrough.

Inputs, all written by frontend/e2e/capture-manual.mjs:

    screenshots/*.png   the pictures
    results.json        what was typed and what the dashboard computed
    contoh-brief/       the policy briefs the dashboard generated

No result is typed into the builder. The only numbers written by hand are the
lecturer's instructor key (manual/data.py, from sheet 11 of
DeciBridge_Demo_Training_Package.xlsx). Every result in the manual is read from
results.json and compared with that key; if a deterministic or BIA figure
disagrees, the build stops rather than print a manual that is wrong.

The PDF is printed from an HTML rendering by the Chromium that Playwright
installed, not exported from Word: automating Word attaches to whatever Word
window the user has open, and closes it when done.

    backend/.venv/Scripts/python.exe docs/panduan/build_manual.py
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

from docx import Document

from manual.appendices import appendix_differences, appendix_key, appendix_problems
from manual.chapters_analysis import about, bia, cover, ekonomi, preparation, psa, ringkasan
from manual.chapters_decision import brief, etd, rekomendasi, signoff, versi
from manual.data import (
    CASE_ID, FOOTER, HERE, INITIAL, MISMATCHES, OUT_DOCX, OUT_PDF, PRINTER,
)
from manual.document import Manual


def check_brief() -> None:
    """The brief table in chapter 9 claims these figures are in the generated file."""
    path = HERE / "contoh-brief" / f"{CASE_ID}_policy_brief_v1.docx"
    if not path.exists():
        sys.exit(f"Brief missing: {path.name}. Run the capture first.")
    document = Document(str(path))
    text = "\n".join(
        [p.text for p in document.paragraphs]
        + [c.text for t in document.tables for r in t.rows for c in r.cells]
    )
    det = INITIAL["deterministic"]
    expected = [
        f"Rp {Decimal(det['icer']):,.0f}".replace(",", "."),
        f"{Decimal(det['total_qaly_intervention']):.4f}",
        f"{Decimal(det['total_qaly_comparator']):.4f}",
        f"Rp {Decimal(INITIAL['bia']['cumulative_net_impact']):,.0f}".replace(",", "."),
        "ADOPSI BERSYARAT",
    ]
    for needle in expected:
        if needle not in text:
            MISMATCHES.append(f"Brief v1 tidak memuat '{needle}'")


def build() -> None:
    check_brief()
    with tempfile.TemporaryDirectory() as workdir:
        m = Manual(Path(workdir))
        cover(m)
        about(m)
        preparation(m)
        ringkasan(m)
        ekonomi(m)
        bia(m)
        psa(m)
        etd(m)
        rekomendasi(m)
        signoff(m)
        brief(m)
        versi(m)
        appendix_key(m)
        appendix_differences(m)
        appendix_problems(m)

        if MISMATCHES:
            sys.exit("Manual NOT built; the dashboard disagrees with the key:\n  " + "\n  ".join(MISMATCHES))

        title = "Buku Panduan DeciBridge - Kasus Latihan ARNI vs Ramipril"
        m.doc.core_properties.title = title
        m.doc.core_properties.subject = f"Kasus {CASE_ID}, data latihan sintetis"
        try:
            m.doc.save(str(OUT_DOCX))
        except PermissionError:
            sys.exit(f"{OUT_DOCX.name} is open in Word. Close it and run this again.")
        megabytes = OUT_DOCX.stat().st_size / 1_048_576
        print(f"{OUT_DOCX.name}: {m.figures} gambar, {m.tables} tabel, {megabytes:.1f} MB")

        page = Path(workdir) / "manual.html"
        page.write_text(m.html_document(title), encoding="utf-8")
        printed = subprocess.run(
            ["node", str(PRINTER), str(page), str(OUT_PDF), FOOTER],
            capture_output=True, text=True, check=False,
        )
        if printed.returncode != 0:
            sys.exit(f"PDF not produced:\n{printed.stderr.strip()}")
        print(f"{OUT_PDF.name}: {OUT_PDF.stat().st_size / 1_048_576:.1f} MB")


if __name__ == "__main__":
    build()
