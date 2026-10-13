"""Writes the manual twice from one set of calls: a DOCX to edit, and HTML to print as PDF."""

from __future__ import annotations

import re
import sys
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps

from .data import FOOTER, SHOTS

TEXT_WIDTH_CM = 17.0
NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


_MARKUP = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _runs(paragraph, text: str, *, size: float | None = None, color: RGBColor | None = None,
          italic: bool = False) -> None:
    """Adds text with **bold** for screen labels and `mono` for typed values."""
    for part in _MARKUP.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        run.italic = italic or None


def _html(text: str) -> str:
    """The same **bold** / `mono` markup, as HTML."""
    out = []
    for part in _MARKUP.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append(f"<strong>{escape(part[2:-2])}</strong>")
        elif part.startswith("`"):
            out.append(f"<code>{escape(part[1:-1])}</code>")
        else:
            out.append(escape(part))
    return "".join(out)


CSS = """
@page { size: A4; margin: 18mm 18mm 20mm 18mm; }
* { box-sizing: border-box; }
body { font-family: Calibri, Carlito, Arial, sans-serif; font-size: 10.5pt; line-height: 1.35;
       color: #1a1a1a; margin: 0; }
h1 { font-size: 17pt; color: #1F4E79; margin: 0 0 8pt; break-before: page; }
h2 { font-size: 13pt; color: #1F4E79; margin: 14pt 0 6pt; break-after: avoid; }
h3 { font-size: 11pt; color: #1F4E79; margin: 10pt 0 4pt; break-after: avoid; }
p { margin: 0 0 6pt; }
code { font-family: Consolas, "Courier New", monospace; font-size: 9.5pt; background: #f1f3f5;
       padding: 0 2pt; border-radius: 2pt; }
ul, ol { margin: 0 0 6pt; padding-left: 16pt; }
li { margin-bottom: 2pt; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 10pt; font-size: 9pt; }
th { background: #1F4E79; color: #fff; text-align: left; padding: 3pt 5pt; border: 0.5pt solid #1F4E79; }
td { padding: 3pt 5pt; border: 0.5pt solid #8a8a8a; vertical-align: top; }
tr { break-inside: avoid; }
thead { display: table-header-group; }
.caption-table { font-size: 9.5pt; margin: 8pt 0 2pt; break-after: avoid; }
.role { background: #EAF1F8; padding: 4pt 7pt; margin: 0 0 8pt; font-size: 10pt; }
.box { padding: 6pt 8pt; margin: 8pt 0; font-size: 10pt; break-inside: avoid; }
.box p { margin: 0 0 3pt; }
.note { background: #EAF1F8; } .warn { background: #FFF3D6; } .check { background: #E6F4EA; }
figure { margin: 8pt 0 12pt; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; border: 0.5pt solid #c8c8c8; }
figcaption { font-size: 9pt; color: #595959; margin-top: 3pt; }
.cover { height: 250mm; display: flex; flex-direction: column; justify-content: center;
         text-align: center; }
.cover .title { font-size: 28pt; font-weight: 700; color: #1F4E79; }
.cover .subtitle { font-size: 16pt; color: #1F4E79; margin-top: 6pt; }
.cover .lead { font-size: 12pt; color: #595959; margin: 6pt 0 40pt; }
.cover .meta { font-size: 11pt; margin: 1pt 0; }
.cover .box { text-align: left; margin-top: 40pt; }
"""


class Manual:
    """Writes the same content twice: a DOCX to edit and an HTML page to print as PDF."""

    def __init__(self, workdir: Path) -> None:
        self.doc = Document()
        self.html: list[str] = []
        self.workdir = workdir
        self.figures = 0
        self.tables = 0
        self._chapters = 0
        self._page_setup()

    # layout ---------------------------------------------------------------

    def _page_setup(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)

        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12

        for name, size, before in (("Heading 1", 17, 0), ("Heading 2", 13, 12), ("Heading 3", 11, 8)):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = NAVY
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(footer, f"{FOOTER} - halaman ", size=8, color=GREY)
        self._page_number(footer)

    @staticmethod
    def _page_number(paragraph) -> None:
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            run = paragraph.add_run()
            run.font.size = Pt(8)
            run.font.color.rgb = GREY
            if kind:
                char = OxmlElement("w:fldChar")
                char.set(qn("w:fldCharType"), kind)
                run._r.append(char)
            else:
                instruction = OxmlElement("w:instrText")
                instruction.set(qn("xml:space"), "preserve")
                instruction.text = text
                run._r.append(instruction)

    def _gap(self) -> None:
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(4)

    # text -----------------------------------------------------------------

    def cover(self, title: str, subtitle: str, lead: str, meta: list[str],
              warning_title: str, warning: str) -> None:
        doc = self.doc
        for _ in range(5):
            doc.add_paragraph()
        for text, size, color in ((f"**{title}**", 28, NAVY), (subtitle, 16, NAVY), (lead, 12, GREY)):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(paragraph, text, size=size, color=color)
        for _ in range(3):
            doc.add_paragraph()
        for line in meta:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            _runs(paragraph, line, size=11)
        for _ in range(3):
            doc.add_paragraph()

        self.html.append(
            '<div class="cover">'
            f'<div class="title">{escape(title)}</div>'
            f'<div class="subtitle">{escape(subtitle)}</div>'
            f'<div class="lead">{escape(lead)}</div>'
            + "".join(f'<div class="meta">{_html(line)}</div>' for line in meta)
        )
        self.callout(warning_title, [warning], kind="warn")
        self.html.append("</div>")

    def chapter(self, title: str) -> None:
        self.doc.add_page_break()
        self.doc.add_heading(title, level=1)
        self.html.append(f"<h1>{escape(title)}</h1>")

    def section(self, title: str) -> None:
        self.doc.add_heading(title, level=2)
        self.html.append(f"<h2>{escape(title)}</h2>")

    def text(self, body: str) -> None:
        _runs(self.doc.add_paragraph(), body)
        self.html.append(f"<p>{_html(body)}</p>")

    def bullets(self, items: list[str]) -> None:
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            _runs(paragraph, item)
        self.html.append("<ul>" + "".join(f"<li>{_html(i)}</li>" for i in items) + "</ul>")

    def steps(self, items: list[str]) -> None:
        for index, item in enumerate(items, start=1):
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.75)
            paragraph.paragraph_format.space_after = Pt(3)
            _runs(paragraph, f"{index}.\t{item}")
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
        self.html.append("<ol>" + "".join(f"<li>{_html(i)}</li>" for i in items) + "</ol>")

    def role_line(self, role: str, package_steps: str) -> None:
        line = f"**Peran:** {role}     **Langkah paket:** {package_steps}"
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        _shade(cell, "EAF1F8")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        _runs(paragraph, line, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        self.html.append(f'<div class="role">{_html(line)}</div>')

    def callout(self, title: str, lines: list[str], *, kind: str = "note") -> None:
        fill = {"note": "EAF1F8", "warn": "FFF3D6", "check": "E6F4EA"}[kind]
        self._gap()
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        _shade(cell, fill)
        first = cell.paragraphs[0]
        first.paragraph_format.space_after = Pt(2)
        _runs(first, f"**{title}**", size=10)
        for line in lines:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            _runs(paragraph, line, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        self.html.append(
            f'<div class="box {kind}"><p><strong>{escape(title)}</strong></p>'
            + "".join(f"<p>{_html(line)}</p>" for line in lines)
            + "</div>"
        )

    # tables ---------------------------------------------------------------

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float],
              caption: str | None = None) -> None:
        if caption:
            self.tables += 1
            heading = self.doc.add_paragraph()
            _keep_with_next(heading)
            heading.paragraph_format.space_after = Pt(3)
            _runs(heading, f"**Tabel {self.tables}.** {caption}", size=9.5)
            self.html.append(
                f'<p class="caption-table"><strong>Tabel {self.tables}.</strong> {_html(caption)}</p>'
            )
        else:
            self._gap()

        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        scale = TEXT_WIDTH_CM / sum(widths)

        header = table.rows[0]
        _repeat_header(header)
        for cell, title, width in zip(header.cells, headers, widths):
            cell.width = Cm(width * scale)
            _shade(cell, "1F4E79")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _runs(paragraph, f"**{title}**", size=9, color=RGBColor(0xFF, 0xFF, 0xFF))

        for values in rows:
            cells = table.add_row().cells
            for cell, value, width in zip(cells, values, widths):
                cell.width = Cm(width * scale)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                _runs(paragraph, value, size=9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

        total = sum(widths)
        columns = "".join(f'<col style="width:{100 * w / total:.2f}%">' for w in widths)
        head = "".join(f"<th>{escape(h)}</th>" for h in headers)
        body = "".join(
            "<tr>" + "".join(f"<td>{_html(v)}</td>" for v in values) + "</tr>" for values in rows
        )
        self.html.append(
            f"<table><colgroup>{columns}</colgroup><thead><tr>{head}</tr></thead>"
            f"<tbody>{body}</tbody></table>"
        )

    # pictures -------------------------------------------------------------

    def figure(self, name: str, caption: str, *, max_height_cm: float = 18.5,
               width_cm: float = TEXT_WIDTH_CM, crop_bottom_at: float | None = None,
               crop_top_at: float = 0.0) -> None:
        source = SHOTS / f"{name}.png"
        if not source.exists():
            sys.exit(f"Screenshot missing: {source.name}. Run the capture first.")

        with Image.open(source) as image:
            picture = image.convert("RGB")
        if crop_bottom_at is not None or crop_top_at:
            top = int(picture.height * crop_top_at)
            bottom = int(picture.height * (crop_bottom_at or 1.0))
            picture = picture.crop((0, top, picture.width, bottom))
        if picture.width > 1700:
            ratio = 1700 / picture.width
            picture = picture.resize((1700, int(picture.height * ratio)), Image.LANCZOS)

        self.figures += 1
        prepared = self.workdir / f"{name}-{self.figures}.png"
        bordered = ImageOps.expand(picture, border=2, fill=(200, 200, 200))
        bordered.quantize(colors=256, method=Image.Quantize.MEDIANCUT).save(prepared, optimize=True)

        natural_height = width_cm * picture.height / picture.width
        if natural_height > max_height_cm:
            width_cm = max_height_cm * picture.width / picture.height

        holder = self.doc.add_paragraph()
        holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        holder.paragraph_format.space_after = Pt(2)
        _keep_with_next(holder)
        holder.add_run().add_picture(str(prepared), width=Cm(width_cm))

        label = self.doc.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.paragraph_format.space_after = Pt(10)
        _runs(label, f"**Gambar {self.figures}.** {caption}", size=9, color=GREY)

        self.html.append(
            f'<figure><img src="{prepared.name}" style="width:{width_cm:.2f}cm" alt="">'
            f"<figcaption><strong>Gambar {self.figures}.</strong> {_html(caption)}</figcaption></figure>"
        )

    def html_document(self, title: str) -> str:
        return (
            '<!doctype html><html lang="id"><head><meta charset="utf-8">'
            f"<title>{escape(title)}</title><style>{CSS}</style></head>"
            f"<body>{''.join(self.html)}</body></html>"
        )
